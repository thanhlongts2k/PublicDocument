"""
Module: exam_generator.py
Biên soạn và kết xuất Đề Thi Hoàn Chỉnh (Exam Paper) chuẩn form in ấn PDF Print-Ready.
Tuân thủ nghiêm ngặt quy cách sư phạm, bố cục A4, chống gãy trang và quy tắc Zero-Key.
"""

import os
import json
from datetime import datetime
from docx import Document
from docx.shared import Cm, Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL

from .styles import (
    setup_page_a4, setup_header_footer, format_run, format_paragraph,
    set_cell_background, set_cell_borders, set_cell_margins, set_table_borders_thin,
    COLOR_PRIMARY_NAVY, COLOR_SECONDARY_BLUE, COLOR_BG_HEADER, COLOR_BG_LIGHT,
    COLOR_BORDER_LIGHT, COLOR_TEXT_MAIN, COLOR_TEXT_MUTED, FONT_JAPANESE_VIET
)
from .pdf_converter import convert_docx_to_pdf

class ExamGenerator:
    def __init__(self, data: dict):
        """
        Khởi tạo generator với dữ liệu đề thi dạng dictionary.
        """
        self.data = data
        self.doc = Document()
        setup_page_a4(self.doc)

    def _render_header_and_candidate_box(self):
        """Dựng Phần A: Khung tiêu đề & Khung thông tin thí sinh"""
        info = self.data.get("metadata", {})
        org_name = info.get("organization", "TRUNG TÂM KHẢO THÍ & ĐÀO TẠO NGÔN NGỮ").upper()
        program_name = info.get("program", "CHƯƠNG TRÌNH ĐÀO TẠO TIẾNG NHẬT CHUẨN HOÁ").upper()
        subject_name = info.get("subject", "TIẾNG NHẬT TỔNG HỢP")
        exam_title = info.get("exam_title", "ĐỀ THI ĐÁNH GIÁ NĂNG LỰC").upper()
        exam_code = info.get("exam_code", "101")
        duration = info.get("duration_minutes", 60)
        num_questions = info.get("total_questions", len(self._get_all_questions()))
        exam_date = info.get("date", datetime.now().strftime("%d/%m/%Y"))

        # 1. Header đối xứng 2 cột
        tbl_header = self.doc.add_table(rows=1, cols=2)
        tbl_header.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl_header.autofit = False

        c_left = tbl_header.rows[0].cells[0]
        c_right = tbl_header.rows[0].cells[1]
        c_left.width = Cm(10.0)
        c_right.width = Cm(7.0)

        # Cột trái
        p_org = c_left.paragraphs[0]
        format_paragraph(p_org, space_before_pt=0, space_after_pt=2, align=WD_ALIGN_PARAGRAPH.CENTER)
        r_org = p_org.add_run(org_name)
        format_run(r_org, size_pt=9.5, bold=True, color_hex=COLOR_PRIMARY_NAVY)

        p_prog = c_left.add_paragraph()
        format_paragraph(p_prog, space_before_pt=0, space_after_pt=2, align=WD_ALIGN_PARAGRAPH.CENTER)
        r_prog = p_prog.add_run(program_name)
        format_run(r_prog, size_pt=9.0, bold=False, color_hex=COLOR_TEXT_MUTED)

        p_subj = c_left.add_paragraph()
        format_paragraph(p_subj, space_before_pt=0, space_after_pt=0, align=WD_ALIGN_PARAGRAPH.CENTER)
        r_subj = p_subj.add_run(f"Môn thi: {subject_name}")
        format_run(r_subj, size_pt=9.5, bold=True)

        # Cột phải
        p_title = c_right.paragraphs[0]
        format_paragraph(p_title, space_before_pt=0, space_after_pt=2, align=WD_ALIGN_PARAGRAPH.CENTER)
        r_title = p_title.add_run(exam_title)
        format_run(r_title, size_pt=10.5, bold=True, color_hex=COLOR_PRIMARY_NAVY)

        p_code = c_right.add_paragraph()
        format_paragraph(p_code, space_before_pt=2, space_after_pt=0, align=WD_ALIGN_PARAGRAPH.CENTER)
        r_code_label = p_code.add_run("MÃ ĐỀ THI: ")
        format_run(r_code_label, size_pt=9.5, bold=True)
        r_code_val = p_code.add_run(f" {exam_code} ")
        format_run(r_code_val, size_pt=11.0, bold=True, color_hex=COLOR_PRIMARY_NAVY)

        # Đóng khung viền nổi bật cho mã đề
        set_cell_borders(c_right, 
                         top={'sz': '12', 'color': COLOR_PRIMARY_NAVY},
                         bottom={'sz': '12', 'color': COLOR_PRIMARY_NAVY},
                         left={'sz': '12', 'color': COLOR_PRIMARY_NAVY},
                         right={'sz': '12', 'color': COLOR_PRIMARY_NAVY})
        set_cell_background(c_right, COLOR_BG_LIGHT)
        set_cell_margins(c_right, top_pt=4, bottom_pt=4, left_pt=6, right_pt=6)

        # 2. Dòng thông tin thời gian thi cử
        p_info = self.doc.add_paragraph()
        format_paragraph(p_info, space_before_pt=6, space_after_pt=6, align=WD_ALIGN_PARAGRAPH.CENTER)
        r_info1 = p_info.add_run(f"Thời gian làm bài: {duration} phút (không kể thời gian phát đề) | Số lượng: {num_questions} câu | Ngày thi: {exam_date}")
        format_run(r_info1, size_pt=9.5, italic=True)

        # 3. Khung thí sinh & Chấm thi (Bảng 1 hàng x 3 cột)
        tbl_candidate = self.doc.add_table(rows=1, cols=3)
        tbl_candidate.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl_candidate.autofit = False
        set_table_borders_thin(tbl_candidate, border_color=COLOR_BORDER_LIGHT)

        widths = [Cm(6.8), Cm(5.2), Cm(5.0)]
        for i, w in enumerate(widths):
            tbl_candidate.rows[0].cells[i].width = w
            set_cell_margins(tbl_candidate.rows[0].cells[i], top_pt=4, bottom_pt=4, left_pt=6, right_pt=6)

        c1 = tbl_candidate.rows[0].cells[0]
        p1 = c1.paragraphs[0]
        format_paragraph(p1, space_before_pt=0, space_after_pt=3)
        r1 = p1.add_run("Họ và tên: .................................................\nSBD / Mã HV: ...........................................")
        format_run(r1, size_pt=9.0)

        c2 = tbl_candidate.rows[0].cells[1]
        p2 = c2.paragraphs[0]
        format_paragraph(p2, space_before_pt=0, space_after_pt=3)
        r2 = p2.add_run("Phòng thi / Lớp: ....................................\nChữ ký GT: ..............................................")
        format_run(r2, size_pt=9.0)

        c3 = tbl_candidate.rows[0].cells[2]
        p3 = c3.paragraphs[0]
        format_paragraph(p3, space_before_pt=0, space_after_pt=3)
        r3 = p3.add_run("ĐIỂM: ...................  / Lời phê:\n...................................................................")
        format_run(r3, size_pt=9.0, bold=True)
        set_cell_background(c3, COLOR_BG_LIGHT)

        # Đường kẻ phân cách
        p_div = self.doc.add_paragraph()
        format_paragraph(p_div, space_before_pt=4, space_after_pt=6, align=WD_ALIGN_PARAGRAPH.CENTER)
        r_div = p_div.add_run("―" * 45)
        format_run(r_div, size_pt=10.0, color_hex=COLOR_BORDER_LIGHT)

    def _render_instructions(self):
        """Dựng Phần B: Hướng dẫn làm bài"""
        instructions = self.data.get("instructions", [
            "1. Thí sinh kiểm tra kỹ số trang và tính toàn vẹn của đề thi trước khi bắt đầu làm bài.",
            "2. Mỗi câu hỏi chỉ có DUY NHẤT một đáp án đúng. Hãy đánh dấu hoặc tô đen phương án chọn vào phiếu trả lời.",
            "3. Tuyệt đối không sử dụng tài liệu, từ điển, điện thoại hoặc thiết bị điện tử trong phòng thi."
        ])

        p_inst_title = self.doc.add_paragraph()
        format_paragraph(p_inst_title, space_before_pt=2, space_after_pt=2, keep_with_next=True)
        r_inst_title = p_inst_title.add_run("HƯỚNG DẪN LÀM BÀI:")
        format_run(r_inst_title, size_pt=9.5, bold=True, color_hex=COLOR_PRIMARY_NAVY)

        for line in instructions:
            p_line = self.doc.add_paragraph()
            format_paragraph(p_line, space_before_pt=0, space_after_pt=2, keep_with_next=True)
            r_line = p_line.add_run(line)
            format_run(r_line, size_pt=9.0, italic=True, color_hex=COLOR_TEXT_MUTED)

        # Ngăn cách nhẹ
        p_space = self.doc.add_paragraph()
        format_paragraph(p_space, space_before_pt=0, space_after_pt=4)

    def _render_sections_and_questions(self):
        """Dựng Phần C: Các phần thi & Câu hỏi (Áp dụng layout thông minh & Chống gãy trang)"""
        sections = self.data.get("sections", [])
        q_counter = 1

        for s_idx, section in enumerate(sections):
            sec_title = section.get("title", f"PHẦN {s_idx + 1}").upper()
            sec_desc = section.get("description", "")

            # Tiêu đề Section
            p_sec = self.doc.add_paragraph()
            format_paragraph(p_sec, space_before_pt=8, space_after_pt=2, keep_with_next=True)
            r_sec = p_sec.add_run(sec_title)
            format_run(r_sec, size_pt=10.5, bold=True, color_hex=COLOR_PRIMARY_NAVY)

            if sec_desc:
                p_desc = self.doc.add_paragraph()
                format_paragraph(p_desc, space_before_pt=0, space_after_pt=4, keep_with_next=True)
                r_desc = p_desc.add_run(sec_desc)
                format_run(r_desc, size_pt=9.5, italic=True)

            # Khung Đọc hiểu (Reading Box) nếu có passage
            passage = section.get("reading_passage", "")
            if passage:
                tbl_passage = self.doc.add_table(rows=1, cols=1)
                tbl_passage.alignment = WD_TABLE_ALIGNMENT.CENTER
                tbl_passage.autofit = False
                tbl_passage.rows[0].cells[0].width = Cm(17.0)
                set_cell_background(tbl_passage.rows[0].cells[0], COLOR_BG_LIGHT)
                set_cell_borders(tbl_passage.rows[0].cells[0],
                                 left={'sz': '12', 'color': COLOR_SECONDARY_BLUE},
                                 top={'sz': '4', 'color': COLOR_BORDER_LIGHT},
                                 bottom={'sz': '4', 'color': COLOR_BORDER_LIGHT},
                                 right={'sz': '4', 'color': COLOR_BORDER_LIGHT})
                set_cell_margins(tbl_passage.rows[0].cells[0], top_pt=6, bottom_pt=6, left_pt=8, right_pt=8)

                cell_p = tbl_passage.rows[0].cells[0].paragraphs[0]
                format_paragraph(cell_p, space_before_pt=0, space_after_pt=0, line_spacing=1.2)
                r_pas = cell_p.add_run(passage)
                format_run(r_pas, size_pt=9.5, italic=False)

                # Khoảng trống sau passage
                p_sp = self.doc.add_paragraph()
                format_paragraph(p_sp, space_before_pt=0, space_after_pt=4)

            # Render từng câu hỏi
            questions = section.get("questions", [])
            for q in questions:
                self._render_single_question(q, q_counter)
                q_counter += 1

        # Dòng kết thúc đề thi
        p_end = self.doc.add_paragraph()
        format_paragraph(p_end, space_before_pt=14, space_after_pt=2, align=WD_ALIGN_PARAGRAPH.CENTER, keep_with_next=True)
        r_end = p_end.add_run("―" * 15 + " HẾT " + "―" * 15)
        format_run(r_end, size_pt=10.5, bold=True, color_hex=COLOR_PRIMARY_NAVY)

        p_note = self.doc.add_paragraph()
        format_paragraph(p_note, space_before_pt=0, space_after_pt=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        r_note = p_note.add_run("(Cán bộ coi thi không giải thích gì thêm)")
        format_run(r_note, size_pt=9.0, italic=True, color_hex=COLOR_TEXT_MUTED)

    def _render_single_question(self, q: dict, q_number: int):
        """Render 1 câu hỏi với bố cục trắc nghiệm tự động tối ưu & chống gãy trang"""
        q_text = q.get("question", "")
        q_type = q.get("type", "multiple_choice") # multiple_choice, star_arrangement
        options = q.get("options", []) # Danh sách [ "1. ...", "2. ...", "3. ...", "4. ..." ] hoặc [ "A. ...", ... ]

        # 1. Đoạn text câu hỏi (bật keep_with_next để gắn chặt với đáp án)
        p_q = self.doc.add_paragraph()
        format_paragraph(p_q, space_before_pt=4, space_after_pt=2, keep_with_next=True)
        
        r_num = p_q.add_run(f"Câu {q_number}: ")
        format_run(r_num, size_pt=10.0, bold=True, color_hex=COLOR_PRIMARY_NAVY)
        
        r_txt = p_q.add_run(q_text)
        format_run(r_txt, size_pt=10.0)

        # Xử lý dạng bài ghép câu dấu sao (Star arrangement)
        if q_type == "star_arrangement":
            star_parts = q.get("star_parts", [])
            if star_parts:
                p_star = self.doc.add_paragraph()
                format_paragraph(p_star, space_before_pt=1, space_after_pt=2, keep_with_next=True)
                p_star.paragraph_format.left_indent = Inches(0.2)
                r_star_label = p_star.add_run("Các mảnh ghép:  ")
                format_run(r_star_label, size_pt=9.5, italic=True)
                for part in star_parts:
                    r_p = p_star.add_run(f"   [{part}]   ")
                    format_run(r_p, size_pt=9.5, bold=True)

        if not options:
            return

        # Chuẩn hóa nhãn đáp án (1, 2, 3, 4 hoặc A, B, C, D)
        labels = ["A", "B", "C", "D"]
        normalized_opts = []
        for i, opt in enumerate(options):
            lbl = labels[i] if i < len(labels) else str(i + 1)
            # Nếu opt chưa có tiền tố
            if not opt.startswith(("A.", "B.", "C.", "D.", "1.", "2.", "3.", "4.")):
                normalized_opts.append(f"{lbl}. {opt}")
            else:
                normalized_opts.append(opt)

        # 2. Xác định bố cục phương án: 1 dòng, 2x2, hay 4 dòng
        max_len = max(len(opt) for opt in normalized_opts)

        if max_len <= 15 and len(normalized_opts) == 4:
            # Layout 1 dòng dàn đều
            p_opt = self.doc.add_paragraph()
            format_paragraph(p_opt, space_before_pt=1, space_after_pt=4, keep_with_next=False)
            p_opt.paragraph_format.left_indent = Inches(0.25)
            line_str = "        ".join(normalized_opts)
            r_opt = p_opt.add_run(line_str)
            format_run(r_opt, size_pt=9.5)

        elif max_len <= 35 and len(normalized_opts) == 4:
            # Layout 2 dòng x 2 cột
            p_opt1 = self.doc.add_paragraph()
            format_paragraph(p_opt1, space_before_pt=1, space_after_pt=1, keep_with_next=True)
            p_opt1.paragraph_format.left_indent = Inches(0.25)
            # Căn chỉnh khoảng cách cố định
            spacing = " " * max(2, 35 - len(normalized_opts[0]))
            r_1 = p_opt1.add_run(f"{normalized_opts[0]}{spacing}{normalized_opts[1]}")
            format_run(r_1, size_pt=9.5)

            p_opt2 = self.doc.add_paragraph()
            format_paragraph(p_opt2, space_before_pt=1, space_after_pt=4, keep_with_next=False)
            p_opt2.paragraph_format.left_indent = Inches(0.25)
            spacing2 = " " * max(2, 35 - len(normalized_opts[2]))
            r_2 = p_opt2.add_run(f"{normalized_opts[2]}{spacing2}{normalized_opts[3]}")
            format_run(r_2, size_pt=9.5)

        else:
            # Layout 4 dòng liên tiếp thụt lề
            for idx, opt in enumerate(normalized_opts):
                p_opt = self.doc.add_paragraph()
                is_last = (idx == len(normalized_opts) - 1)
                format_paragraph(p_opt, space_before_pt=1, space_after_pt=3 if is_last else 1, keep_with_next=not is_last)
                p_opt.paragraph_format.left_indent = Inches(0.25)
                r_opt = p_opt.add_run(opt)
                format_run(r_opt, size_pt=9.5)

    def _render_mini_answer_sheet(self):
        """Dựng Phần D: Phiếu tô trả lời trắc nghiệm nhanh (Mini Answer Sheet)"""
        total_q = len(self._get_all_questions())
        if total_q == 0:
            return

        p_sheet_title = self.doc.add_paragraph()
        format_paragraph(p_sheet_title, space_before_pt=10, space_after_pt=4, align=WD_ALIGN_PARAGRAPH.CENTER, keep_with_next=True)
        r_st = p_sheet_title.add_run("PHIẾU TÔ TRẢ LỜI TRẮC NGHIỆM NHANH (ANSWER SHEET)")
        format_run(r_st, size_pt=10.0, bold=True, color_hex=COLOR_PRIMARY_NAVY)

        # Chia theo các khối 10 câu mỗi bảng hoặc bảng lưới nhiều cột
        cols_per_row = 10
        chunk_size = 10
        all_q_ids = list(range(1, total_q + 1))

        # Tạo bảng 2 hàng cho mỗi 10 câu
        for start_idx in range(0, total_q, chunk_size):
            chunk = all_q_ids[start_idx:start_idx + chunk_size]
            actual_cols = len(chunk)

            tbl = self.doc.add_table(rows=2, cols=actual_cols)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            tbl.autofit = False
            set_table_borders_thin(tbl, border_color=COLOR_BORDER_LIGHT)

            col_width = Cm(17.0 / 10)
            for c_i, q_num in enumerate(chunk):
                # Hàng 1: Số câu
                cell_top = tbl.rows[0].cells[c_i]
                cell_top.width = col_width
                set_cell_background(cell_top, COLOR_BG_HEADER)
                set_cell_margins(cell_top, top_pt=2, bottom_pt=2, left_pt=2, right_pt=2)
                p_top = cell_top.paragraphs[0]
                format_paragraph(p_top, space_before_pt=0, space_after_pt=0, align=WD_ALIGN_PARAGRAPH.CENTER)
                r_top = p_top.add_run(f"Câu {q_num}")
                format_run(r_top, size_pt=8.0, bold=True, color_hex=COLOR_PRIMARY_NAVY)

                # Hàng 2: Ô trắc nghiệm [A] [B] [C] [D]
                cell_bot = tbl.rows[1].cells[c_i]
                cell_bot.width = col_width
                set_cell_margins(cell_bot, top_pt=3, bottom_pt=3, left_pt=1, right_pt=1)
                p_bot = cell_bot.paragraphs[0]
                format_paragraph(p_bot, space_before_pt=0, space_after_pt=0, align=WD_ALIGN_PARAGRAPH.CENTER)
                r_bot = p_bot.add_run("[A] [B]\n[C] [D]")
                format_run(r_bot, size_pt=7.5, bold=False, color_hex=COLOR_TEXT_MUTED)

            p_sp = self.doc.add_paragraph()
            format_paragraph(p_sp, space_before_pt=0, space_after_pt=4)

    def _get_all_questions(self):
        questions = []
        for s in self.data.get("sections", []):
            questions.extend(s.get("questions", []))
        return questions

    def generate(self, output_dir: str = ".", keep_docx: bool = False, package: bool = True) -> str:
        """
        Thực thi biên soạn đề thi, xuất file docx tạm thời rồi convert sang PDF Print-Ready.
        Theo quy chuẩn: Đề thi CHỈ TẠO VÀ LƯU FILE PDF.
        Nếu package=True (mặc định), tự động gom file vào thư mục định danh chuẩn: output_dir/[base_filename]/
        Trả về đường dẫn tuyệt đối của file PDF sinh ra.
        """
        info = self.data.get("metadata", {})
        exam_prog = info.get("program_code", "EXAM")
        subject_code = info.get("subject_code", "SUBJ")
        exam_type = info.get("type_code", "DeLuyenTap")
        exam_code = info.get("exam_code", "101")
        date_str = info.get("date_code", datetime.now().strftime("%Y%m%d"))

        # Tên file chuẩn: [Kỳ_Thi/Chương_Trình]_[Môn_Học/Trình_Độ]_[Loại_Đề]_[MaDe]_[YYYYMMDD]
        custom_base = self.data.get("package_name") or info.get("package_name") or info.get("base_filename")
        if custom_base:
            base_filename = "".join(c if c.isalnum() or c in ("-", "_") else "_" for c in custom_base).strip("_")
        else:
            base_filename = f"{exam_prog}_{subject_code}_{exam_type}_{exam_code}_{date_str}"

        # Xác định thư mục lưu trữ: Gói đề thi vào folder riêng biệt theo format chuẩn
        norm_out = os.path.normpath(os.path.abspath(output_dir))
        if package and os.path.basename(norm_out) != base_filename:
            target_dir = os.path.join(norm_out, base_filename)
        else:
            target_dir = norm_out

        os.makedirs(target_dir, exist_ok=True)
        docx_filename = f"{base_filename}.docx"
        docx_path = os.path.join(target_dir, docx_filename)

        # Cài đặt Header / Footer
        header_text = f"{info.get('exam_title', 'ĐỀ THI').upper()} - MÃ ĐỀ {exam_code}"
        setup_header_footer(self.doc, header_text)

        # 1. Phần A: Header & Candidate box
        self._render_header_and_candidate_box()

        # 2. Phần B: Hướng dẫn làm bài
        self._render_instructions()

        # 3. Phần C: Nội dung các phần thi & câu hỏi
        self._render_sections_and_questions()

        # 4. Phần D: Phiếu tô trắc nghiệm nhanh
        self._render_mini_answer_sheet()

        # Lưu file docx
        self.doc.save(docx_path)

        # Tự động lưu vết câu hỏi vào Ngân hàng câu hỏi (Question Bank)
        try:
            from .question_bank import QuestionBank
            qb = QuestionBank()
            qb.ingest_exam(self.data)
        except Exception as e:
            print(f"[WARNING] Không thể lưu vết vào Question Bank: {e}")

        # Convert sang PDF (Quy chuẩn: Đề thi CHỈ tạo file PDF)
        pdf_path = os.path.splitext(docx_path)[0] + ".pdf"
        try:
            pdf_path = convert_docx_to_pdf(docx_path, pdf_path)
            # Dọn dẹp file docx trung gian nếu không yêu cầu giữ
            if not keep_docx and os.path.exists(pdf_path) and os.path.exists(docx_path):
                try:
                    os.remove(docx_path)
                except Exception:
                    pass
        except Exception as e:
            print(f"Cảnh báo khi convert PDF: {e}")
            # Nếu lỗi convert PDF thì giữ lại docx để người dùng có tài liệu
            pdf_path = docx_path

        # Tự động xuất âm thanh MP3 vào cùng thư mục nếu có kịch bản thi nghe đính kèm
        audio_data = self.data.get("audio") or self.data.get("audio_script")
        if audio_data:
            try:
                import asyncio
                from .audio_generator import AudioExamGenerator
                if isinstance(audio_data, dict):
                    if "package_name" not in audio_data:
                        audio_data["package_name"] = base_filename
                    if "metadata" not in audio_data:
                        audio_data["metadata"] = info
                    print(f"[INFO] Phát hiện kịch bản âm thanh đính kèm! Đang xuất file MP3 vào cùng thư mục gói...")
                    audio_gen = AudioExamGenerator(audio_data)
                    asyncio.run(audio_gen.generate(output_dir=target_dir, package=False))
            except Exception as e:
                print(f"[WARNING] Không thể tự động tạo âm thanh đính kèm: {e}")

        return pdf_path
