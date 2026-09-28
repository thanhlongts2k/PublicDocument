"""
Module: styles.py
Định nghĩa hệ thống style, typography, căn lề và các hàm XML helper cho python-docx
Tuân thủ chuẩn sư phạm & in ấn chuyên nghiệp (Word & PDF Print-Ready).
"""

from docx import Document
from docx.shared import Cm, Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

# --- BẢNG MÀU CHUẨN THIẾT KẾ ---
COLOR_PRIMARY_NAVY = "1F4E79"      # Xanh navy chuẩn học thuật
COLOR_SECONDARY_BLUE = "2E5B82"    # Xanh dương phụ trợ
COLOR_ACCENT_ORANGE = "D96B27"     # Cam nhấn mạnh
COLOR_CORRECT_GREEN = "1E7E34"     # Xanh lá [ĐÚNG]
COLOR_WRONG_RED = "BD2130"         # Đỏ [CHƯA ĐÚNG]
COLOR_BG_HEADER = "E9EEF4"         # Nền header bảng xám navy nhạt
COLOR_BG_LIGHT = "F8F9FA"          # Nền thẻ box / reading box
COLOR_BG_CALLOUT = "F0F4F8"        # Nền khung Callout
COLOR_BORDER_LIGHT = "CBD5E1"      # Viền mỏng
COLOR_TEXT_MAIN = "1A1A1A"         # Chữ chính (đen than)
COLOR_TEXT_MUTED = "555555"        # Chữ phụ mờ

FONT_JAPANESE_VIET = "BIZ UDPGothic" # Font cực nét cho song ngữ Nhật - Việt
FONT_BACKUP = "Meiryo"
FONT_ENGLISH_LATIN = "Times New Roman"

def setup_page_a4(doc: Document, top_cm=2.0, bottom_cm=2.0, left_cm=2.5, right_cm=1.5):
    """Thiết lập khổ giấy A4 và lề chuẩn in ấn & đóng gáy"""
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(top_cm)
        section.bottom_margin = Cm(bottom_cm)
        section.left_margin = Cm(left_cm)   # 2.5cm để bấm kim / đóng gáy
        section.right_margin = Cm(right_cm) # 1.5cm
        section.header_distance = Cm(1.0)
        section.footer_distance = Cm(1.0)

def set_cell_background(cell, hex_color: str):
    """Tô màu nền cho ô bảng Word (Shading)"""
    tcPr = cell._tc.get_or_add_tcPr()
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shading_elm)

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    """
    Đặt viền tùy chỉnh cho ô bảng.
    Mỗi tham số là dict ví dụ: {'val': 'single', 'sz': '4', 'color': '1F4E79'}
    """
    tcPr = cell._tc.get_or_add_tcPr()
    top_xml = f'<w:top w:val="{top.get("val", "single")}" w:sz="{top.get("sz", "4")}" w:space="0" w:color="{top.get("color", "auto")}"/>' if top else '<w:top w:val="none"/>'
    left_xml = f'<w:left w:val="{left.get("val", "single")}" w:sz="{left.get("sz", "4")}" w:space="0" w:color="{left.get("color", "auto")}"/>' if left else '<w:left w:val="none"/>'
    bottom_xml = f'<w:bottom w:val="{bottom.get("val", "single")}" w:sz="{bottom.get("sz", "4")}" w:space="0" w:color="{bottom.get("color", "auto")}"/>' if bottom else '<w:bottom w:val="none"/>'
    right_xml = f'<w:right w:val="{right.get("val", "single")}" w:sz="{right.get("sz", "4")}" w:space="0" w:color="{right.get("color", "auto")}"/>' if right else '<w:right w:val="none"/>'
    
    borders_xml = f'<w:tcBorders {nsdecls("w")}>{top_xml}{left_xml}{bottom_xml}{right_xml}</w:tcBorders>'
    tcPr.append(parse_xml(borders_xml))

def set_cell_margins(cell, top_pt=4, bottom_pt=4, left_pt=6, right_pt=6):
    """Đặt padding bên trong ô (dxa = 1/20 pt)"""
    tcPr = cell._tc.get_or_add_tcPr()
    top_dxa = int(top_pt * 20)
    bot_dxa = int(bottom_pt * 20)
    left_dxa = int(left_pt * 20)
    right_dxa = int(right_pt * 20)
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top_dxa}" w:type="dxa"/>
            <w:bottom w:w="{bot_dxa}" w:type="dxa"/>
            <w:left w:w="{left_dxa}" w:type="dxa"/>
            <w:right w:w="{right_dxa}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders_thin(table, border_color="CBD5E1"):
    """Đặt viền mỏng toàn bộ cho bảng"""
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>
            <w:insideV w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def format_run(run, font_name=FONT_JAPANESE_VIET, size_pt=10.5, bold=False, italic=False, color_hex=None):
    """Định dạng chuẩn cho 1 đoạn chữ (run)"""
    run.font.name = font_name
    # Đảm bảo font chữ hỗ trợ ký tự Á Đông (Kanji, Kana)
    rPr = run._r.get_or_add_rPr()
    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{font_name}" w:hAnsi="{font_name}" w:eastAsia="{font_name}" w:cs="{font_name}"/>')
    rPr.append(rFonts)

    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color_hex:
        r, g, b = int(color_hex[0:2], 16), int(color_hex[2:4], 16), int(color_hex[4:6], 16)
        run.font.color.rgb = RGBColor(r, g, b)

def format_paragraph(p, space_before_pt=0, space_after_pt=4, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT, keep_with_next=False):
    """Thiết lập khoảng cách dòng và đoạn văn bản"""
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before_pt)
    p.paragraph_format.space_after = Pt(space_after_pt)
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.keep_with_next = keep_with_next

def add_page_number_to_run(run):
    """Chèn field số trang hiện tại (PAGE)"""
    fld1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instr = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fld2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fld3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)
    run._r.append(fld3)

def add_numpages_to_run(run):
    """Chèn field tổng số trang (NUMPAGES)"""
    fld1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instr = parse_xml(r'<w:instrText %s xml:space="preserve"> NUMPAGES </w:instrText>' % nsdecls('w'))
    fld2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fld3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)
    run._r.append(fld3)

def setup_header_footer(doc: Document, header_text: str):
    """Thiết lập Header mã đề và Footer Trang X / Y chuẩn form"""
    for section in doc.sections:
        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.text = ""
        format_paragraph(hp, space_before_pt=0, space_after_pt=0, align=WD_ALIGN_PARAGRAPH.RIGHT)
        hrun = hp.add_run(header_text)
        format_run(hrun, size_pt=8.5, color_hex=COLOR_TEXT_MUTED)

        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.text = ""
        format_paragraph(fp, space_before_pt=4, space_after_pt=0, align=WD_ALIGN_PARAGRAPH.CENTER)
        
        frun1 = fp.add_run("Trang ")
        format_run(frun1, size_pt=9.0, color_hex=COLOR_TEXT_MUTED)
        frun_page = fp.add_run()
        format_run(frun_page, size_pt=9.0, color_hex=COLOR_TEXT_MUTED)
        add_page_number_to_run(frun_page)
        
        frun2 = fp.add_run(" / ")
        format_run(frun2, size_pt=9.0, color_hex=COLOR_TEXT_MUTED)
        
        frun_numpages = fp.add_run()
        format_run(frun_numpages, size_pt=9.0, color_hex=COLOR_TEXT_MUTED)
        add_numpages_to_run(frun_numpages)
