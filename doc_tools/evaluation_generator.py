"""
Module: evaluation_generator.py
Biên soạn và kết xuất "Báo Cáo Đánh Giá, Phân Tích Lỗi Sai & Giải Thích Chi Tiết"
Chuẩn Microsoft Word và PDF, mang tính sư phạm và học thuật cao.
"""

import os
import json
from datetime import datetime
from docx import Document
from docx.shared import Cm, Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL

from .styles import (
    setup_page_a4, setup_header_footer, format_run, format_paragraph,
    set_cell_background, set_cell_borders, set_cell_margins, set_table_borders_thin,
    COLOR_PRIMARY_NAVY, COLOR_SECONDARY_BLUE, COLOR_CORRECT_GREEN, COLOR_WRONG_RED,
    COLOR_BG_HEADER, COLOR_BG_LIGHT, COLOR_BG_CALLOUT, COLOR_BORDER_LIGHT,
    COLOR_TEXT_MAIN, COLOR_TEXT_MUTED, FONT_JAPANESE_VIET
)
from .pdf_converter import convert_docx_to_pdf

class EvaluationGenerator:
    def __init__(self, data: dict):
        self.data = data
        self.doc = Document()
        setup_page_a4(self.doc)

    def _render_header_and_metadata(self):
        """Dựng Phần 1: Header & Thẻ thông tin định danh (Metadata Box 2x4)"""
        meta = self.data.get("metadata", {})
        title = meta.get("title", "BÁO CÁO ĐÁNH GIÁ, PHÂN TÍCH LỖI SAI & GIẢI THÍCH CHI TIẾT").upper()
        subtitle = meta.get("subtitle", "Khung Đào Tạo Nâng Cao Năng Lực Học Thuật & Chiến Thuật Phòng Thi")
        student_name = meta.get("student_name", "Nguyễn Văn A")
        date_str = meta.get("date", datetime.now().strftime("%d/%m/%Y"))
        exam_set = meta.get("exam_set", "Bộ đề kiểm tra định kỳ")
        score_summary = meta.get("score_summary", "12/15 (80%)")
        core_strengths = meta.get("core_strengths", "Nắm chắc từ vựng nền tảng, phản xạ nhanh với các dạng câu hỏi trực diện.")
        key_weaknesses = meta.get("key_weaknesses", "Dễ sập bẫy các cấu trúc ngữ pháp sắc thái tương đồng (がち vs 気味) và trợ từ liên kết.")

        # Tiêu đề chính
        p_title = self.doc.add_paragraph()
        format_paragraph(p_title, space_before_pt=0, space_after_pt=2, align=WD_ALIGN_PARAGRAPH.CENTER, keep_with_next=True)
        r_title = p_title.add_run(title)
        format_run(r_title, size_pt=13.0, bold=True, color_hex=COLOR_PRIMARY_NAVY)

        # Phụ đề
        p_sub = self.doc.add_paragraph()
        format_paragraph(p_sub, space_before_pt=0, space_after_pt=6, align=WD_ALIGN_PARAGRAPH.CENTER, keep_with_next=True)
        r_sub = p_sub.add_run(subtitle)
        format_run(r_sub, size_pt=9.5, italic=True, color_hex=COLOR_TEXT_MUTED)

        # Bảng Metadata Card 2 hàng x 4 cột
        tbl = self.doc.add_table(rows=2, cols=4)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        set_table_borders_thin(tbl, border_color=COLOR_BORDER_LIGHT)

        # Hàng 1
        headers_r1 = [
            ("Học viên", student_name),
            ("Ngày thực hiện", date_str),
            ("Bộ đề / Bài thi", exam_set),
            ("Tổng kết quả", score_summary)
        ]
        widths_r1 = [Cm(4.0), Cm(4.0), Cm(5.0), Cm(4.0)]

        for c_i, (label, val) in enumerate(headers_r1):
            cell = tbl.rows[0].cells[c_i]
            cell.width = widths_r1[c_i]
            set_cell_background(cell, COLOR_BG_HEADER)
            set_cell_margins(cell, top_pt=4, bottom_pt=4, left_pt=6, right_pt=6)
            p = cell.paragraphs[0]
            format_paragraph(p, space_before_pt=0, space_after_pt=1, align=WD_ALIGN_PARAGRAPH.CENTER)
            r_lbl = p.add_run(f"{label}\n")
            format_run(r_lbl, size_pt=8.5, bold=True, color_hex=COLOR_PRIMARY_NAVY)
            r_val = p.add_run(val)
            format_run(r_val, size_pt=9.5, bold=(c_i == 3), color_hex=COLOR_PRIMARY_NAVY if c_i == 3 else COLOR_TEXT_MAIN)

        # Hàng 2: Gộp ô: 2 ô đầu cho Điểm mạnh, 2 ô sau cho Lỗ hổng
        # Merge cell 0 và 1
        cell_str = tbl.rows[1].cells[0]
        cell_str.merge(tbl.rows[1].cells[1])
        cell_str.width = Cm(8.0)
        set_cell_margins(cell_str, top_pt=4, bottom_pt=4, left_pt=6, right_pt=6)
        p_str = cell_str.paragraphs[0]
        format_paragraph(p_str, space_before_pt=0, space_after_pt=1)
        r_str_lbl = p_str.add_run("✔ Điểm mạnh cốt lõi: ")
        format_run(r_str_lbl, size_pt=8.5, bold=True, color_hex=COLOR_CORRECT_GREEN)
        r_str_val = p_str.add_run(core_strengths)
        format_run(r_str_val, size_pt=9.0)

        # Merge cell 2 và 3 (sau khi merge 0-1, cell 2 và 3 vẫn truy cập qua tbl.rows[1].cells[2])
        cell_weak = tbl.rows[1].cells[2]
        cell_weak.merge(tbl.rows[1].cells[3])
        cell_weak.width = Cm(9.0)
        set_cell_margins(cell_weak, top_pt=4, bottom_pt=4, left_pt=6, right_pt=6)
        p_weak = cell_weak.paragraphs[0]
        format_paragraph(p_weak, space_before_pt=0, space_after_pt=1)
        r_weak_lbl = p_weak.add_run("⚠ Lỗ hổng cần xử lý ngay: ")
        format_run(r_weak_lbl, size_pt=8.5, bold=True, color_hex=COLOR_WRONG_RED)
        r_weak_val = p_weak.add_run(key_weaknesses)
        format_run(r_weak_val, size_pt=9.0)

        p_space = self.doc.add_paragraph()
        format_paragraph(p_space, space_before_pt=4, space_after_pt=4)

    def _render_summary_table(self):
        """Dựng Phần 2: Nội dung chi tiết & Bảng kết quả tổng hợp"""
        results = self.data.get("summary_results", [])
        if not results:
            return

        p_h2 = self.doc.add_paragraph()
        format_paragraph(p_h2, space_before_pt=6, space_after_pt=4, keep_with_next=True)
        r_h2 = p_h2.add_run("PHẦN I: BẢNG KẾT QUẢ TỔNG HỢP & GIẢI THÍCH CHI TIẾT")
        format_run(r_h2, size_pt=10.5, bold=True, color_hex=COLOR_PRIMARY_NAVY)

        tbl = self.doc.add_table(rows=1 + len(results), cols=4)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        set_table_borders_thin(tbl, border_color=COLOR_BORDER_LIGHT)

        widths = [Cm(1.2), Cm(2.2), Cm(2.2), Cm(11.4)]
        headers = ["STT", "Lựa chọn", "Đáp án", "Nội dung câu gốc, Bản dịch & Giải thích cô đọng"]

        # Render Header
        for c_i, h in enumerate(headers):
            cell = tbl.rows[0].cells[c_i]
            cell.width = widths[c_i]
            set_cell_background(cell, COLOR_PRIMARY_NAVY)
            set_cell_margins(cell, top_pt=4, bottom_pt=4, left_pt=4, right_pt=4)
            p = cell.paragraphs[0]
            align = WD_ALIGN_PARAGRAPH.CENTER if c_i < 3 else WD_ALIGN_PARAGRAPH.LEFT
            format_paragraph(p, space_before_pt=0, space_after_pt=0, align=align)
            r = p.add_run(h)
            format_run(r, size_pt=9.0, bold=True, color_hex="FFFFFF")

        # Render rows
        for r_i, item in enumerate(results):
            row_cells = tbl.rows[r_i + 1].cells
            for c_i in range(4):
                row_cells[c_i].width = widths[c_i]
                set_cell_margins(row_cells[c_i], top_pt=3, bottom_pt=3, left_pt=4, right_pt=4)

            # Cột 1: Câu số
            p0 = row_cells[0].paragraphs[0]
            format_paragraph(p0, space_before_pt=0, space_after_pt=0, align=WD_ALIGN_PARAGRAPH.CENTER)
            r0 = p0.add_run(str(item.get("question_number", r_i + 1)))
            format_run(r0, size_pt=9.0, bold=True)

            # Cột 2: Học viên chọn
            p1 = row_cells[1].paragraphs[0]
            format_paragraph(p1, space_before_pt=0, space_after_pt=0, align=WD_ALIGN_PARAGRAPH.CENTER)
            r1 = p1.add_run(str(item.get("student_choice", "")))
            is_correct = item.get("is_correct", False)
            format_run(r1, size_pt=9.0, bold=True, color_hex=COLOR_CORRECT_GREEN if is_correct else COLOR_WRONG_RED)

            # Cột 3: Đáp án chuẩn
            p2 = row_cells[2].paragraphs[0]
            format_paragraph(p2, space_before_pt=0, space_after_pt=0, align=WD_ALIGN_PARAGRAPH.CENTER)
            r2 = p2.add_run(str(item.get("correct_answer", "")))
            format_run(r2, size_pt=9.0, bold=True, color_hex=COLOR_PRIMARY_NAVY)

            # Cột 4: Nội dung & Giải thích
            p3 = row_cells[3].paragraphs[0]
            format_paragraph(p3, space_before_pt=0, space_after_pt=0)

            # Gắn nhãn [ĐÚNG] hoặc [CHƯA ĐÚNG]
            r_tag = p3.add_run("[ĐÚNG] " if is_correct else "[CHƯA ĐÚNG] ")
            format_run(r_tag, size_pt=8.5, bold=True, color_hex=COLOR_CORRECT_GREEN if is_correct else COLOR_WRONG_RED)

            r_orig = p3.add_run(item.get("question_content", "") + "\n")
            format_run(r_orig, size_pt=9.0, bold=True)

            trans = item.get("translation", "")
            if trans:
                r_trans = p3.add_run(f"Dịch: {trans}\n")
                format_run(r_trans, size_pt=8.5, italic=True, color_hex=COLOR_TEXT_MUTED)

            expl = item.get("brief_explanation", "")
            if expl:
                r_expl = p3.add_run(f"Giải thích: {expl}")
                format_run(r_expl, size_pt=8.5)

            # Tô màu nền xen kẽ nhẹ
            if r_i % 2 == 1:
                for c_i in range(4):
                    set_cell_background(row_cells[c_i], COLOR_BG_LIGHT)

        p_space = self.doc.add_paragraph()
        format_paragraph(p_space, space_before_pt=4, space_after_pt=4)

    def _render_error_analysis(self):
        """Dựng Phần 3: Phân tích chuyên sâu lỗi sai & Tư duy nguyên nhân (Quy trình 3 bước)"""
        errors = self.data.get("error_analyses", [])
        if not errors:
            return

        p_h3 = self.doc.add_paragraph()
        format_paragraph(p_h3, space_before_pt=8, space_after_pt=4, keep_with_next=True)
        r_h3 = p_h3.add_run("PHẦN II: PHÂN TÍCH CHUYÊN SÂU LỖI SAI & TƯ DUY NGUYÊN NHÂN")
        format_run(r_h3, size_pt=10.5, bold=True, color_hex=COLOR_PRIMARY_NAVY)

        for err in errors:
            q_num = err.get("question_number", "")
            q_text = err.get("question_text", "")
            phenomenon = err.get("error_phenomenon", "")
            root_cause = err.get("root_cause", "")
            core_rule = err.get("core_rule", "")

            # Tiêu đề câu sai
            p_q = self.doc.add_paragraph()
            format_paragraph(p_q, space_before_pt=6, space_after_pt=2, keep_with_next=True)
            r_num = p_q.add_run(f"◆ CÂU SỐ {q_num}: ")
            format_run(r_num, size_pt=9.5, bold=True, color_hex=COLOR_WRONG_RED)
            r_qt = p_q.add_run(q_text)
            format_run(r_qt, size_pt=9.5, bold=True)

            # Bước 1: Hiện tượng sai
            p_step1 = self.doc.add_paragraph()
            format_paragraph(p_step1, space_before_pt=1, space_after_pt=2, keep_with_next=True)
            p_step1.paragraph_format.left_indent = Inches(0.15)
            r_s1_lbl = p_step1.add_run("1. Đề bài & Hiện tượng sai: ")
            format_run(r_s1_lbl, size_pt=9.0, bold=True, color_hex=COLOR_PRIMARY_NAVY)
            r_s1_txt = p_step1.add_run(phenomenon)
            format_run(r_s1_txt, size_pt=9.0)

            # Bước 2: Nguồn gốc nhầm lẫn
            p_step2 = self.doc.add_paragraph()
            format_paragraph(p_step2, space_before_pt=1, space_after_pt=3, keep_with_next=True)
            p_step2.paragraph_format.left_indent = Inches(0.15)
            r_s2_lbl = p_step2.add_run("2. Nguồn gốc nhầm lẫn: ")
            format_run(r_s2_lbl, size_pt=9.0, bold=True, color_hex=COLOR_PRIMARY_NAVY)
            r_s2_txt = p_step2.add_run(root_cause)
            format_run(r_s2_txt, size_pt=9.0)

            # Bước 3: Khung Callout Quote Quy tắc cốt lõi
            tbl_callout = self.doc.add_table(rows=1, cols=1)
            tbl_callout.alignment = WD_TABLE_ALIGNMENT.CENTER
            tbl_callout.autofit = False
            tbl_callout.rows[0].cells[0].width = Cm(16.5)
            set_cell_background(tbl_callout.rows[0].cells[0], COLOR_BG_CALLOUT)
            set_cell_borders(tbl_callout.rows[0].cells[0],
                             left={'sz': '18', 'color': COLOR_PRIMARY_NAVY},
                             top=None, bottom=None, right=None)
            set_cell_margins(tbl_callout.rows[0].cells[0], top_pt=5, bottom_pt=5, left_pt=10, right_pt=8)

            cp = tbl_callout.rows[0].cells[0].paragraphs[0]
            format_paragraph(cp, space_before_pt=0, space_after_pt=0)
            r_c_lbl = cp.add_run("★ QUY TẮC CỐT LÕI (BẢN CHẤT KIẾN THỨC):\n")
            format_run(r_c_lbl, size_pt=8.5, bold=True, color_hex=COLOR_PRIMARY_NAVY)
            r_c_txt = cp.add_run(core_rule)
            format_run(r_c_txt, size_pt=9.0, italic=False)

            p_sp = self.doc.add_paragraph()
            format_paragraph(p_sp, space_before_pt=0, space_after_pt=3)

    def _render_confusion_matrix_and_traps(self):
        """Dựng Phần 4: Ma trận phân biệt cấu trúc tương tự (Confusion Matrix) & Bóc trần cạm bẫy đề thi"""
        matrix_data = self.data.get("confusion_matrix", {})
        if not matrix_data:
            return

        p_h4 = self.doc.add_paragraph()
        format_paragraph(p_h4, space_before_pt=8, space_after_pt=4, keep_with_next=True)
        r_h4 = p_h4.add_run("PHẦN III: MA TRẬN PHÂN BIỆT CẤU TRÚC DỄ NHẦM LẪN & BẪY ĐỀ THI")
        format_run(r_h4, size_pt=10.5, bold=True, color_hex=COLOR_PRIMARY_NAVY)

        # 1. Bảng Confusion Matrix
        rows_data = matrix_data.get("rows", [])
        if rows_data:
            p_tbl_lbl = self.doc.add_paragraph()
            format_paragraph(p_tbl_lbl, space_before_pt=2, space_after_pt=2, keep_with_next=True)
            r_lbl = p_tbl_lbl.add_run("1. Bảng đối chiếu các cấu trúc / kiến thức tương tự (Confusion Matrix):")
            format_run(r_lbl, size_pt=9.5, bold=True)

            tbl = self.doc.add_table(rows=1 + len(rows_data), cols=5)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            tbl.autofit = False
            set_table_borders_thin(tbl, border_color=COLOR_BORDER_LIGHT)

            headers = ["Cấu trúc / Điểm ngữ pháp", "Ý nghĩa cốt lõi", "Điều kiện kết hợp", "Sắc thái & Ngữ cảnh", "Ví dụ phân biệt thực tế"]
            widths = [Cm(3.0), Cm(3.2), Cm(3.0), Cm(3.6), Cm(4.2)]

            for c_i, h in enumerate(headers):
                cell = tbl.rows[0].cells[c_i]
                cell.width = widths[c_i]
                set_cell_background(cell, COLOR_SECONDARY_BLUE)
                set_cell_margins(cell, top_pt=4, bottom_pt=4, left_pt=4, right_pt=4)
                p = cell.paragraphs[0]
                format_paragraph(p, space_before_pt=0, space_after_pt=0, align=WD_ALIGN_PARAGRAPH.CENTER)
                r = p.add_run(h)
                format_run(r, size_pt=8.5, bold=True, color_hex="FFFFFF")

            for r_i, row in enumerate(rows_data):
                row_cells = tbl.rows[r_i + 1].cells
                for c_i in range(5):
                    row_cells[c_i].width = widths[c_i]
                    set_cell_margins(row_cells[c_i], top_pt=3, bottom_pt=3, left_pt=4, right_pt=4)
                    if r_i % 2 == 1:
                        set_cell_background(row_cells[c_i], COLOR_BG_LIGHT)

                # Col 0: Cấu trúc
                p0 = row_cells[0].paragraphs[0]
                format_paragraph(p0, space_before_pt=0, space_after_pt=0)
                r0 = p0.add_run(row.get("structure", ""))
                format_run(r0, size_pt=8.5, bold=True, color_hex=COLOR_PRIMARY_NAVY)

                # Col 1: Ý nghĩa
                p1 = row_cells[1].paragraphs[0]
                format_paragraph(p1, space_before_pt=0, space_after_pt=0)
                r1 = p1.add_run(row.get("meaning", ""))
                format_run(r1, size_pt=8.5)

                # Col 2: Kết hợp
                p2 = row_cells[2].paragraphs[0]
                format_paragraph(p2, space_before_pt=0, space_after_pt=0)
                r2 = p2.add_run(row.get("connection", ""))
                format_run(r2, size_pt=8.5, italic=True)

                # Col 3: Sắc thái
                p3 = row_cells[3].paragraphs[0]
                format_paragraph(p3, space_before_pt=0, space_after_pt=0)
                r3 = p3.add_run(row.get("nuance", ""))
                format_run(r3, size_pt=8.5)

                # Col 4: Ví dụ
                p4 = row_cells[4].paragraphs[0]
                format_paragraph(p4, space_before_pt=0, space_after_pt=0)
                r4 = p4.add_run(row.get("example", ""))
                format_run(r4, size_pt=8.5)

            p_sp = self.doc.add_paragraph()
            format_paragraph(p_sp, space_before_pt=0, space_after_pt=4)

        # 2. Cạm bẫy đề thi
        traps = matrix_data.get("exam_traps", [])
        if traps:
            p_trap_lbl = self.doc.add_paragraph()
            format_paragraph(p_trap_lbl, space_before_pt=4, space_after_pt=2, keep_with_next=True)
            r_t_lbl = p_trap_lbl.add_run("2. Bóc trần cạm bẫy đề thi (Exam Traps & Distractors):")
            format_run(r_t_lbl, size_pt=9.5, bold=True, color_hex=COLOR_PRIMARY_NAVY)

            for trap in traps:
                p_tr = self.doc.add_paragraph()
                format_paragraph(p_tr, space_before_pt=1, space_after_pt=2)
                p_tr.paragraph_format.left_indent = Inches(0.2)
                r_dot = p_tr.add_run("▸ ")
                format_run(r_dot, size_pt=9.0, bold=True, color_hex=COLOR_PRIMARY_NAVY)
                r_name = p_tr.add_run(trap.get("trap_name", "") + ": ")
                format_run(r_name, size_pt=9.0, bold=True)
                r_desc = p_tr.add_run(trap.get("trap_description", ""))
                format_run(r_desc, size_pt=9.0)

        # 3. Mẹo kiểm tra phản xạ nhanh (3 giây)
        checklist = matrix_data.get("quick_checklist", [])
        if checklist:
            p_chk_lbl = self.doc.add_paragraph()
            format_paragraph(p_chk_lbl, space_before_pt=4, space_after_pt=2, keep_with_next=True)
            r_c_lbl = p_chk_lbl.add_run("3. Mẹo kiểm tra phản xạ nhanh (Quick Checklist trong 3 giây):")
            format_run(r_c_lbl, size_pt=9.5, bold=True, color_hex=COLOR_PRIMARY_NAVY)

            for tip in checklist:
                p_tip = self.doc.add_paragraph()
                format_paragraph(p_tip, space_before_pt=1, space_after_pt=2)
                p_tip.paragraph_format.left_indent = Inches(0.2)
                r_chk = p_tip.add_run("⚡ ")
                format_run(r_chk, size_pt=9.0, bold=True, color_hex=COLOR_SECONDARY_BLUE)
                r_txt = p_tip.add_run(tip)
                format_run(r_txt, size_pt=9.0)

    def _render_drill_exercises(self):
        """Dựng Phần 5: Bài tập bù đắp bẫy câu hỏi (Drill Exercises - Không kèm đáp án)"""
        exercises = self.data.get("drill_exercises", [])
        if not exercises:
            return

        p_h5 = self.doc.add_paragraph()
        format_paragraph(p_h5, space_before_pt=8, space_after_pt=2, keep_with_next=True)
        r_h5 = p_h5.add_run("PHẦN IV: BÀI TẬP BÙ ĐẮP BẪY CÂU HỎI (DRILL EXERCISES)")
        format_run(r_h5, size_pt=10.5, bold=True, color_hex=COLOR_PRIMARY_NAVY)

        p_note = self.doc.add_paragraph()
        format_paragraph(p_note, space_before_pt=0, space_after_pt=4, keep_with_next=True)
        r_note = p_note.add_run("*(Học viên tự làm các câu hỏi sau để củng cố phản xạ trước các bẫy đề thi vừa phân tích)*")
        format_run(r_note, size_pt=9.0, italic=True, color_hex=COLOR_TEXT_MUTED)

        for idx, ex in enumerate(exercises):
            p_ex = self.doc.add_paragraph()
            format_paragraph(p_ex, space_before_pt=4, space_after_pt=2, keep_with_next=True)
            r_num = p_ex.add_run(f"Câu {idx + 1}: ")
            format_run(r_num, size_pt=9.5, bold=True, color_hex=COLOR_PRIMARY_NAVY)
            r_txt = p_ex.add_run(ex.get("question", ""))
            format_run(r_txt, size_pt=9.5)

            opts = ex.get("options", [])
            if opts:
                p_opt = self.doc.add_paragraph()
                format_paragraph(p_opt, space_before_pt=1, space_after_pt=3)
                p_opt.paragraph_format.left_indent = Inches(0.25)
                # Dàn đều hoặc chia dòng
                opts_str = "       ".join(opts)
                r_opt = p_opt.add_run(opts_str)
                format_run(r_opt, size_pt=9.0)

    def generate(self, output_dir: str = ".", export_pdf: bool = False) -> str:
        """
        Thực thi xuất file đánh giá.
        Theo quy chuẩn: File ĐÁNH GIÁ CHỈ TẠO FILE WORD (.docx).
        Trả về đường dẫn tuyệt đối của file docx.
        """
        meta = self.data.get("metadata", {})
        sub_code = meta.get("subject_code", "JLPT_N3")
        doc_type = meta.get("type_code", "DanhGia")
        student_code = meta.get("student_code", "HocVienA")
        topic_code = meta.get("topic_code", "NguPhapTuan01")
        date_str = meta.get("date_code", datetime.now().strftime("%Y%m%d"))
        version = meta.get("version", "1")

        # Tên file: [Mã_Môn]_[Loại_Tài_Liệu]_[Tên_Học_Viên]_[Nội_Dung]_[YYYYMMDD]_v[PhiênBản].docx
        base_filename = f"{sub_code}_{doc_type}_{student_code}_{topic_code}_{date_str}_v{version}"
        docx_filename = f"{base_filename}.docx"
        docx_path = os.path.join(output_dir, docx_filename)

        # Header / Footer
        header_text = f"BÁO CÁO ĐÁNH GIÁ NĂNG LỰC - {student_code.upper()} - {date_str}"
        setup_header_footer(self.doc, header_text)

        # Render 5 sections
        self._render_header_and_metadata()
        self._render_summary_table()
        self._render_error_analysis()
        self._render_confusion_matrix_and_traps()
        self._render_drill_exercises()

        # Lưu file docx
        os.makedirs(output_dir, exist_ok=True)
        self.doc.save(docx_path)

        # Tự động lưu vết lỗi sai và bài tập bù đắp vào Ngân hàng câu hỏi (Question Bank)
        try:
            from .question_bank import QuestionBank
            qb = QuestionBank()
            qb.ingest_evaluation(self.data)
        except Exception as e:
            print(f"[WARNING] Không thể lưu vết đánh giá vào Question Bank: {e}")

        # Chỉ convert sang PDF nếu có cờ export_pdf rõ ràng
        if export_pdf:
            pdf_path = os.path.splitext(docx_path)[0] + ".pdf"
            try:
                convert_docx_to_pdf(docx_path, pdf_path)
            except Exception as e:
                print(f"Cảnh báo khi convert PDF: {e}")

        return docx_path
